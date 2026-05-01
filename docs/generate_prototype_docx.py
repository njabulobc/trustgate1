from pathlib import Path
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING

SRC = Path('docs/trustgate_prototype_document.md')
OUT = Path('docs/trustgate_prototype_document.docx')

text = SRC.read_text(encoding='utf-8')

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

for raw in text.splitlines():
    line = raw.rstrip()
    if not line:
        doc.add_paragraph('')
        continue
    if line.startswith('# '):
        p = doc.add_paragraph(line[2:])
        p.style = doc.styles['Heading 1']
        continue
    if re.match(r'^##\s+', line):
        p = doc.add_paragraph(line[3:])
        p.style = doc.styles['Heading 2']
        continue
    if re.match(r'^###\s+', line):
        p = doc.add_paragraph(line[4:])
        p.style = doc.styles['Heading 3']
        continue
    if line.startswith('```'):
        doc.add_paragraph(line)
        continue
    if line.startswith('- ') or re.match(r'^\d+\.\s+', line):
        doc.add_paragraph(line)
        continue
    doc.add_paragraph(line)

doc.save(OUT)
print(f'Wrote {OUT}')
