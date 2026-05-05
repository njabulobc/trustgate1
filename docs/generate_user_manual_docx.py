from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "trustgate_user_manual.md"
OUT_DIR = ROOT / "output" / "doc"
OUT = OUT_DIR / "TrustGate_User_Manual.docx"


def set_cell_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)

    for name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True

    if "Code Block" not in document.styles:
        style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Consolas"
        style.font.size = Pt(9.5)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1


def add_title_page(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TrustGate MVP User Manual")
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.bold = True

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Operational guide for onboarding, screening, risk, escalation, monitoring, and reporting")
    run.font.name = "Aptos"
    run.font.size = Pt(11)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Generated from the current repository implementation")
    run.italic = True

    document.add_paragraph("")
    document.add_page_break()


def add_code_paragraph(document: Document, line: str) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.add_run(line)
    set_cell_shading(paragraph, "F3F4F6")


def add_body(document: Document, text: str) -> None:
    in_code_block = False
    for raw in text.splitlines():
        line = raw.rstrip()

        if line.startswith("```"):
            in_code_block = not in_code_block
            continue

        if in_code_block:
            add_code_paragraph(document, line)
            continue

        if not line:
            document.add_paragraph("")
            continue

        if line.startswith("# "):
            document.add_paragraph(line[2:], style="Heading 1")
            continue

        if line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 2")
            continue

        if line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 3")
            continue

        if re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(re.sub(r"^\d+\.\s+", "", line))
            continue

        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:])
            continue

        document.add_paragraph(line)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    text = SRC.read_text(encoding="utf-8")

    document = Document()
    configure_styles(document)
    add_title_page(document)
    add_body(document, text)
    document.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
